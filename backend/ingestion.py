#!/usr/bin/env python3
# ═══════════════════════════════════════════════════════════════════════════
# Document Ingestion Utilities
# Text extraction and chunking for PDF and XLSX documents
# ═══════════════════════════════════════════════════════════════════════════

import io
import re
from dataclasses import dataclass
from typing import Optional

import csv
import pdfplumber
import openpyxl
import docx


# ───────────────────────────────────────────────────────────────────────────────
# Types
# ───────────────────────────────────────────────────────────────────────────────

@dataclass
class PageText:
    page_num: int  # 1-indexed page number
    text:     str


@dataclass
class PdfChunk:
    text:       str
    page_start: int  # First page this chunk appears on (1-indexed)
    page_end:   int  # Last page this chunk appears on (1-indexed)


# ───────────────────────────────────────────────────────────────────────────────
# PDF Text Extraction
# ───────────────────────────────────────────────────────────────────────────────

def extract_pdf_text(pdf_bytes: bytes) -> str:
    """Extract text content from a PDF file.

    Args:
        pdf_bytes: Raw bytes of the PDF file

    Returns:
        Extracted text with pages separated by double newlines
    """
    text_parts = []

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                # Clean up common PDF extraction artifacts
                page_text = clean_pdf_text(page_text)
                text_parts.append(page_text)

    return "\n\n".join(text_parts)


def extract_pdf_pages(pdf_bytes: bytes) -> list[PageText]:
    """Extract text content from a PDF file with page numbers.

    Args:
        pdf_bytes: Raw bytes of the PDF file

    Returns:
        List of PageText objects with page numbers (1-indexed)
    """
    pages = []

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text()
            if page_text:
                page_text = clean_pdf_text(page_text)
                pages.append(PageText(page_num=i + 1, text=page_text))

    return pages


def clean_pdf_text(text: str) -> str:
    """Clean common artifacts from PDF-extracted text.

    Args:
        text: Raw extracted text from PDF

    Returns:
        Cleaned text
    """
    # ─────────────────────────────────────────────────────────────────
    # Fix common ligature issues
    # ─────────────────────────────────────────────────────────────────
    ligatures = {
        'ﬁ': 'fi',
        'ﬂ': 'fl',
        'ﬀ': 'ff',
        'ﬃ': 'ffi',
        'ﬄ': 'ffl',
    }
    for lig, replacement in ligatures.items():
        text = text.replace(lig, replacement)

    # ─────────────────────────────────────────────────────────────────
    # Normalize whitespace
    # ─────────────────────────────────────────────────────────────────
    # Replace multiple spaces with single space
    text = re.sub(r'[ \t]+', ' ', text)

    # Normalize line endings
    text = text.replace('\r\n', '\n').replace('\r', '\n')

    # Remove excessive blank lines (more than 2 consecutive)
    text = re.sub(r'\n{3,}', '\n\n', text)

    # ─────────────────────────────────────────────────────────────────
    # Fix hyphenation at line breaks
    # ─────────────────────────────────────────────────────────────────
    # Join words split across lines with hyphen
    text = re.sub(r'-\n([a-z])', r'\1', text)

    return text.strip()


# ───────────────────────────────────────────────────────────────────────────────
# Text Chunking
# ───────────────────────────────────────────────────────────────────────────────

def chunk_text(
    text:         str,
    chunk_size:   int = 1000,
    chunk_overlap: int = 200
) -> list[str]:
    """Split text into overlapping chunks.

    Uses paragraph boundaries when possible, falling back to
    sentence boundaries for large paragraphs.

    Args:
        text:          Text to chunk
        chunk_size:    Target size for each chunk in characters
        chunk_overlap: Number of characters to overlap between chunks

    Returns:
        List of text chunks
    """
    if not text.strip():
        return []

    # If text is smaller than chunk size, return as single chunk
    if len(text) <= chunk_size:
        return [text.strip()]

    chunks = []

    # ─────────────────────────────────────────────────────────────────
    # Split into paragraphs first
    # ─────────────────────────────────────────────────────────────────
    paragraphs = re.split(r'\n\n+', text)

    current_chunk = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue

        # ─────────────────────────────────────────────────────────────
        # Handle paragraphs larger than chunk size
        # ─────────────────────────────────────────────────────────────
        if len(para) > chunk_size:
            # Save current chunk if exists
            if current_chunk:
                chunks.append(current_chunk.strip())
                current_chunk = ""

            # Split large paragraph by sentences
            sentences = split_into_sentences(para)
            for sentence in sentences:
                if len(current_chunk) + len(sentence) + 1 > chunk_size:
                    if current_chunk:
                        chunks.append(current_chunk.strip())
                        # Start new chunk with overlap from previous
                        current_chunk = get_overlap_text(current_chunk, chunk_overlap)
                    current_chunk += (" " if current_chunk else "") + sentence
                else:
                    current_chunk += (" " if current_chunk else "") + sentence

        # ─────────────────────────────────────────────────────────────
        # Normal paragraph handling
        # ─────────────────────────────────────────────────────────────
        elif len(current_chunk) + len(para) + 2 > chunk_size:
            # Adding this paragraph would exceed limit
            chunks.append(current_chunk.strip())
            # Start new chunk with overlap
            current_chunk = get_overlap_text(current_chunk, chunk_overlap) + "\n\n" + para
        else:
            current_chunk += ("\n\n" if current_chunk else "") + para

    # Don't forget the last chunk
    if current_chunk.strip():
        chunks.append(current_chunk.strip())

    return chunks


def split_into_sentences(text: str) -> list[str]:
    """Split text into sentences.

    Args:
        text: Text to split

    Returns:
        List of sentences
    """
    # Split on sentence-ending punctuation followed by space or newline
    sentences = re.split(r'(?<=[.!?])\s+', text)
    return [s.strip() for s in sentences if s.strip()]


def get_overlap_text(text: str, overlap_size: int) -> str:
    """Get the last `overlap_size` characters from text, breaking at word boundary.

    Args:
        text:         Source text
        overlap_size: Target overlap size in characters

    Returns:
        Overlap text starting at a word boundary
    """
    if len(text) <= overlap_size:
        return text

    # Take last `overlap_size` characters
    overlap = text[-overlap_size:]

    # Find first word boundary (space) and start from there
    first_space = overlap.find(' ')
    if first_space > 0:
        overlap = overlap[first_space + 1:]

    return overlap.strip()


# ───────────────────────────────────────────────────────────────────────────────
# Page-Aware Chunking
# ───────────────────────────────────────────────────────────────────────────────

def chunk_pdf_with_pages(
    pages:         list[PageText],
    chunk_size:    int = 1000,
    chunk_overlap: int = 200
) -> list[PdfChunk]:
    """Split PDF pages into chunks while tracking page numbers.

    Args:
        pages:         List of PageText objects from extract_pdf_pages
        chunk_size:    Target size for each chunk in characters
        chunk_overlap: Number of characters to overlap between chunks

    Returns:
        List of PdfChunk objects with page range information
    """
    if not pages:
        return []

    chunks = []
    current_text = ""
    current_page_start = pages[0].page_num
    current_page_end = pages[0].page_num

    for page in pages:
        page_text = page.text.strip()
        if not page_text:
            continue

        # ─────────────────────────────────────────────────────────────────
        # Check if adding this page exceeds chunk size
        # ─────────────────────────────────────────────────────────────────
        if len(current_text) + len(page_text) + 2 > chunk_size:
            # Save current chunk if exists
            if current_text:
                chunks.append(PdfChunk(
                    text=current_text.strip(),
                    page_start=current_page_start,
                    page_end=current_page_end
                ))

            # Start new chunk with overlap
            overlap = get_overlap_text(current_text, chunk_overlap)
            current_text = overlap + ("\n\n" if overlap else "") + page_text
            current_page_start = page.page_num
            current_page_end = page.page_num
        else:
            # Add to current chunk
            current_text += ("\n\n" if current_text else "") + page_text
            current_page_end = page.page_num

    # Don't forget the last chunk
    if current_text.strip():
        chunks.append(PdfChunk(
            text=current_text.strip(),
            page_start=current_page_start,
            page_end=current_page_end
        ))

    # ─────────────────────────────────────────────────────────────────────
    # Handle chunks that are still too large (split within page)
    # ─────────────────────────────────────────────────────────────────────
    final_chunks = []
    for chunk in chunks:
        if len(chunk.text) <= chunk_size:
            final_chunks.append(chunk)
        else:
            # Split large chunk while preserving page info
            sub_texts = split_large_text(chunk.text, chunk_size, chunk_overlap)
            for sub_text in sub_texts:
                final_chunks.append(PdfChunk(
                    text=sub_text,
                    page_start=chunk.page_start,
                    page_end=chunk.page_end
                ))

    return final_chunks


def split_large_text(
    text:          str,
    chunk_size:    int,
    chunk_overlap: int
) -> list[str]:
    """Split large text into smaller chunks with overlap."""
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    sentences = split_into_sentences(text)
    current = ""

    for sentence in sentences:
        if len(current) + len(sentence) + 1 > chunk_size:
            if current:
                chunks.append(current.strip())
                current = get_overlap_text(current, chunk_overlap)
            current += (" " if current else "") + sentence
        else:
            current += (" " if current else "") + sentence

    if current.strip():
        chunks.append(current.strip())

    return chunks


# ═══════════════════════════════════════════════════════════════════════════════
# XLSX Ingestion
# ═══════════════════════════════════════════════════════════════════════════════

@dataclass
class XlsxRow:
    """Represents a single row from an XLSX file."""
    sheet_name: str
    row_num:    int      # 1-indexed row number
    text:       str      # Row content as "Header1: Value1, Header2: Value2, ..."
    headers:    list[str]
    values:     list[str]


@dataclass
class XlsxChunk:
    """Represents a chunk of rows from an XLSX file."""
    text:        str
    sheet_name:  str
    row_start:   int  # First row in chunk (1-indexed)
    row_end:     int  # Last row in chunk (1-indexed)


def extract_xlsx_rows(xlsx_bytes: bytes) -> list[XlsxRow]:
    """Extract rows from an XLSX file with header context.

    Each row is converted to a text representation using column headers.
    First row of each sheet is assumed to be headers.

    Args:
        xlsx_bytes: Raw bytes of the XLSX file

    Returns:
        List of XlsxRow objects
    """
    rows = []

    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        headers = []
        first_row = True

        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            # ─────────────────────────────────────────────────────────────
            # Skip completely empty rows
            # ─────────────────────────────────────────────────────────────
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            # ─────────────────────────────────────────────────────────────
            # First non-empty row is headers
            # ─────────────────────────────────────────────────────────────
            if first_row:
                headers = [str(cell).strip() if cell is not None else f"Column_{i+1}"
                          for i, cell in enumerate(row)]
                first_row = False
                continue

            # ─────────────────────────────────────────────────────────────
            # Convert row to text with header context
            # ─────────────────────────────────────────────────────────────
            values = []
            parts = []

            for i, cell in enumerate(row):
                if cell is not None:
                    cell_str = str(cell).strip()
                    if cell_str:
                        header = headers[i] if i < len(headers) else f"Column_{i+1}"
                        parts.append(f"{header}: {cell_str}")
                        values.append(cell_str)
                    else:
                        values.append("")
                else:
                    values.append("")

            if parts:  # Only add rows with actual content
                rows.append(XlsxRow(
                    sheet_name=sheet_name,
                    row_num=row_idx,
                    text=", ".join(parts),
                    headers=headers,
                    values=values
                ))

    workbook.close()
    return rows


def chunk_xlsx_rows(
    rows:           list[XlsxRow],
    rows_per_chunk: int = 10,
    max_chunk_size: int = 2000
) -> list[XlsxChunk]:
    """Group XLSX rows into chunks for embedding.

    Args:
        rows:           List of XlsxRow objects from extract_xlsx_rows
        rows_per_chunk: Target number of rows per chunk
        max_chunk_size: Maximum characters per chunk

    Returns:
        List of XlsxChunk objects
    """
    if not rows:
        return []

    chunks = []
    current_texts = []
    current_sheet = rows[0].sheet_name
    current_row_start = rows[0].row_num
    current_row_end = rows[0].row_num
    current_size = 0

    for row in rows:
        row_text = row.text

        # ─────────────────────────────────────────────────────────────────
        # Sheet change forces new chunk
        # ─────────────────────────────────────────────────────────────────
        if row.sheet_name != current_sheet:
            if current_texts:
                chunks.append(XlsxChunk(
                    text="\n".join(current_texts),
                    sheet_name=current_sheet,
                    row_start=current_row_start,
                    row_end=current_row_end
                ))
            current_texts = [row_text]
            current_sheet = row.sheet_name
            current_row_start = row.row_num
            current_row_end = row.row_num
            current_size = len(row_text)
            continue

        # ─────────────────────────────────────────────────────────────────
        # Check if adding this row exceeds limits
        # ─────────────────────────────────────────────────────────────────
        new_size = current_size + len(row_text) + 1  # +1 for newline
        would_exceed_rows = len(current_texts) >= rows_per_chunk
        would_exceed_size = new_size > max_chunk_size

        if current_texts and (would_exceed_rows or would_exceed_size):
            # Save current chunk
            chunks.append(XlsxChunk(
                text="\n".join(current_texts),
                sheet_name=current_sheet,
                row_start=current_row_start,
                row_end=current_row_end
            ))
            # Start new chunk
            current_texts = [row_text]
            current_row_start = row.row_num
            current_row_end = row.row_num
            current_size = len(row_text)
        else:
            # Add to current chunk
            current_texts.append(row_text)
            current_row_end = row.row_num
            current_size = new_size

    # ─────────────────────────────────────────────────────────────────────
    # Don't forget the last chunk
    # ─────────────────────────────────────────────────────────────────────
    if current_texts:
        chunks.append(XlsxChunk(
            text="\n".join(current_texts),
            sheet_name=current_sheet,
            row_start=current_row_start,
            row_end=current_row_end
        ))

    return chunks


# ═══════════════════════════════════════════════════════════════════════════════
# CSV Ingestion
# ═══════════════════════════════════════════════════════════════════════════════

@dataclass
class CsvRow:
    """Represents a single row from a CSV file."""
    row_num:  int      # 1-indexed row number
    text:     str      # Row content as "Header1: Value1, Header2: Value2, ..."
    headers:  list[str]
    values:   list[str]


@dataclass
class CsvChunk:
    """Represents a chunk of rows from a CSV file."""
    text:      str
    row_start: int  # First row in chunk (1-indexed)
    row_end:   int  # Last row in chunk (1-indexed)


def extract_csv_rows(csv_bytes: bytes) -> list[CsvRow]:
    """Extract rows from a CSV file with header context.

    Each row is converted to a text representation using column headers.
    First row is assumed to be headers.

    Args:
        csv_bytes: Raw bytes of the CSV file

    Returns:
        List of CsvRow objects
    """
    rows = []

    # Decode bytes - try UTF-8 first, fall back to latin-1
    try:
        content = csv_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = csv_bytes.decode("latin-1")

    # Use csv.Sniffer to detect delimiter
    try:
        dialect = csv.Sniffer().sniff(content[:4096])
    except csv.Error:
        dialect = csv.excel  # Default to comma-separated

    reader = csv.reader(io.StringIO(content), dialect)
    headers = []
    first_row = True

    for row_idx, row in enumerate(reader, start=1):
        # ─────────────────────────────────────────────────────────────────
        # Skip completely empty rows
        # ─────────────────────────────────────────────────────────────────
        if all(cell.strip() == "" for cell in row):
            continue

        # ─────────────────────────────────────────────────────────────────
        # First non-empty row is headers
        # ─────────────────────────────────────────────────────────────────
        if first_row:
            headers = [cell.strip() if cell.strip() else f"Column_{i+1}"
                      for i, cell in enumerate(row)]
            first_row = False
            continue

        # ─────────────────────────────────────────────────────────────────
        # Convert row to text with header context
        # ─────────────────────────────────────────────────────────────────
        values = []
        parts = []

        for i, cell in enumerate(row):
            cell_str = cell.strip()
            if cell_str:
                header = headers[i] if i < len(headers) else f"Column_{i+1}"
                parts.append(f"{header}: {cell_str}")
                values.append(cell_str)
            else:
                values.append("")

        if parts:  # Only add rows with actual content
            rows.append(CsvRow(
                row_num=row_idx,
                text=", ".join(parts),
                headers=headers,
                values=values
            ))

    return rows


def chunk_csv_rows(
    rows:           list[CsvRow],
    rows_per_chunk: int = 10,
    max_chunk_size: int = 2000
) -> list[CsvChunk]:
    """Group CSV rows into chunks for embedding.

    Args:
        rows:           List of CsvRow objects from extract_csv_rows
        rows_per_chunk: Target number of rows per chunk
        max_chunk_size: Maximum characters per chunk

    Returns:
        List of CsvChunk objects
    """
    if not rows:
        return []

    chunks = []
    current_texts = []
    current_row_start = rows[0].row_num
    current_row_end = rows[0].row_num
    current_size = 0

    for row in rows:
        row_text = row.text

        # ─────────────────────────────────────────────────────────────────
        # Check if adding this row exceeds limits
        # ─────────────────────────────────────────────────────────────────
        new_size = current_size + len(row_text) + 1  # +1 for newline
        would_exceed_rows = len(current_texts) >= rows_per_chunk
        would_exceed_size = new_size > max_chunk_size

        if current_texts and (would_exceed_rows or would_exceed_size):
            # Save current chunk
            chunks.append(CsvChunk(
                text="\n".join(current_texts),
                row_start=current_row_start,
                row_end=current_row_end
            ))
            # Start new chunk
            current_texts = [row_text]
            current_row_start = row.row_num
            current_row_end = row.row_num
            current_size = len(row_text)
        else:
            # Add to current chunk
            current_texts.append(row_text)
            current_row_end = row.row_num
            current_size = new_size

    # ─────────────────────────────────────────────────────────────────────
    # Don't forget the last chunk
    # ─────────────────────────────────────────────────────────────────────
    if current_texts:
        chunks.append(CsvChunk(
            text="\n".join(current_texts),
            row_start=current_row_start,
            row_end=current_row_end
        ))

    return chunks


# ═══════════════════════════════════════════════════════════════════════════════
# DOCX Ingestion
# ═══════════════════════════════════════════════════════════════════════════════

@dataclass
class DocxParagraph:
    """Represents a paragraph from a DOCX file."""
    text:       str
    style_name: str    # e.g., "Heading 1", "Normal", "List Paragraph"
    is_heading: bool


@dataclass
class DocxChunk:
    """Represents a chunk from a DOCX file."""
    text:         str
    heading_path: str   # e.g., "Introduction > Overview"


def extract_docx_paragraphs(docx_bytes: bytes) -> list[DocxParagraph]:
    """Extract paragraphs from a DOCX file.

    Args:
        docx_bytes: Raw bytes of the DOCX file

    Returns:
        List of DocxParagraph objects
    """
    paragraphs = []

    document = docx.Document(io.BytesIO(docx_bytes))

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = para.style.name if para.style else "Normal"
        is_heading = style_name.startswith("Heading")

        paragraphs.append(DocxParagraph(
            text=text,
            style_name=style_name,
            is_heading=is_heading
        ))

    return paragraphs


def chunk_docx_paragraphs(
    paragraphs:     list[DocxParagraph],
    chunk_size:     int = 1500,
    chunk_overlap:  int = 200
) -> list[DocxChunk]:
    """Chunk DOCX paragraphs, preserving heading context.

    Args:
        paragraphs:    List of DocxParagraph objects
        chunk_size:    Target chunk size in characters
        chunk_overlap: Character overlap between chunks

    Returns:
        List of DocxChunk objects
    """
    if not paragraphs:
        return []

    chunks = []
    current_texts = []
    current_size = 0
    heading_stack = []  # Track heading hierarchy

    def get_heading_path() -> str:
        return " > ".join(heading_stack) if heading_stack else ""

    def flush_chunk():
        nonlocal current_texts, current_size
        if current_texts:
            chunks.append(DocxChunk(
                text="\n\n".join(current_texts),
                heading_path=get_heading_path()
            ))
            current_texts = []
            current_size = 0

    for para in paragraphs:
        # ─────────────────────────────────────────────────────────────────
        # Handle headings - update heading stack
        # ─────────────────────────────────────────────────────────────────
        if para.is_heading:
            # Flush current chunk before new heading
            flush_chunk()

            # Parse heading level (e.g., "Heading 1" -> 1)
            try:
                level = int(para.style_name.split()[-1])
            except (ValueError, IndexError):
                level = 1

            # Adjust heading stack to current level
            while len(heading_stack) >= level:
                heading_stack.pop()
            heading_stack.append(para.text)

            # Add heading as start of new chunk
            current_texts.append(para.text)
            current_size = len(para.text)
            continue

        # ─────────────────────────────────────────────────────────────────
        # Regular paragraph - check chunk size
        # ─────────────────────────────────────────────────────────────────
        para_len = len(para.text)

        if current_size + para_len + 2 > chunk_size and current_texts:
            # Would exceed limit - flush and start new chunk
            flush_chunk()

            # Add overlap from previous content if available
            if chunks:
                overlap = get_overlap_text(chunks[-1].text, chunk_overlap)
                if overlap:
                    current_texts.append(overlap)
                    current_size = len(overlap)

        current_texts.append(para.text)
        current_size += para_len + 2  # +2 for paragraph separator

    # Don't forget the last chunk
    flush_chunk()

    return chunks
